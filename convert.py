#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
批量把 Excel 报表里的金额列（元）换算成万元。
用法示例：
    python convert.py              # 默认：报表练习文件夹，保留 2 位小数
    python convert.py -d 1         # 保留 1 位小数
    python convert.py -d 0         # 保留 0 位（整数）
    python convert.py --folder "D:/某路径" -d 2
"""

import os              # 操作系统接口：拼接路径、创建文件夹、列出目录
import sys             # 系统相关：用于出错时退出程序（可选）
import time            # 计时模块：记录每个文件处理耗时
import argparse        # 命令行参数解析：用 -d 控制保留几位小数
from decimal import Decimal, ROUND_HALF_UP  # 高精度小数 + 四舍五入模式
import openpyxl        # 读写 .xlsx /.xlsm，且不破坏原表格样式


# ---------- 配置项（按需修改）----------
# 列名里只要包含下面任意一个关键字，这一列就会被换算
KEYWORDS = ['本期金额', '年初余额', '期末余额']

# 小数位 -> 对应的 Decimal 量化单位（保留几位就量化到 10 的负 n 次方）
# quantize 后数值会被“卡”到这个精度，例如 2 位就是精确到 0.01
QUANT = {
    2: Decimal('0.01'),   # 保留 2 位小数
    1: Decimal('0.1'),    # 保留 1 位小数
    0: Decimal('1'),      # 保留 0 位（整数）
}

DIVISOR = Decimal('10000')   # 1 万元 = 10000 元，换算除数


def find_header_row(sheet, keywords):
    """扫描工作表，返回表头所在行号（从 1 开始）；找不到返回 None"""
    for row in sheet.iter_rows():            # iter_rows()：逐行遍历整张表
        for cell in row:                     # 取出这一行里的每个单元格
            # cell.value 是单元格内容；str() 防 None；any() 判断是否有任一关键字命中
            if cell.value and any(k in str(cell.value) for k in keywords):
                return cell.row              # cell.row 是行号（从 1 开始）
    return None


def to_wan(value, quant):
    """
    把单个单元格的值（单位：元）换算成万元，并四舍五入到指定小数位。
    返回 Decimal（openpyxl 能直接当作数字存进去，不会丢精度）。
    """
    if value is None:
        return None                          # 空单元格：原样跳过

    if isinstance(value, str):               # 字符串类型
        if value.startswith('='):           # 以等号开头 = 公式单元格，不改动
            return value
        # 去掉常见干扰字符：千分位逗号、货币符号、单位“元”
        cleaned = value.replace(',', '').replace('¥', '').replace('元', '').strip()
        if cleaned == '':                    # 清掉后啥也不剩：当空处理
            return None
        d = Decimal(cleaned)                 # 字符串 -> Decimal
    else:
        # int / float 先转成字符串再进 Decimal，避免 float 本身的二进制误差
        d = Decimal(str(value))

    # 元 -> 万元：除以 10000，再按 quant 量化（四舍五入，逢五进一）
    result = (d / DIVISOR).quantize(quant, rounding=ROUND_HALF_UP)
    return result


def convert_file(path, decimals):
    """换算单个 Excel 文件，返回处理了多少数据行"""
    # load_workbook 默认加载方式会保留字体、颜色、边框、列宽等格式
    wb = openpyxl.load_workbook(path)
    quant = QUANT[decimals]                  # 取出本次要用的精度单位
    total_rows = 0                          # 累计处理的数据行数

    for ws in wb.worksheets:                 # 遍历工作簿里的每个工作表
        hrow = find_header_row(ws, KEYWORDS)  # 找表头行
        if hrow is None:
            continue                        # 这一页没有目标列，跳过

        # 收集需要换算的列号（列号从 1 开始）
        target_cols = []
        for cell in ws[hrow]:               # ws[行号]：取这一整行的单元格
            if cell.value and any(k in str(cell.value) for k in KEYWORDS):
                target_cols.append(cell.column)   # cell.column：列号

        if not target_cols:                 # 没找到可换算列，跳过
            continue

        # 从表头下一行开始，到最大行结束
        for r in range(hrow + 1, ws.max_row + 1):
            row_touched = False             # 这一行有没有被改动过
            for c in target_cols:
                cell = ws.cell(row=r, column=c)   # 精确定位单元格
                new_val = to_wan(cell.value, quant)
                # 有值、且换算后和原来不一样，才写回
                if new_val is not None and new_val != cell.value:
                    cell.value = new_val    # 只改数值，不动样式 -> 格式保留
                    row_touched = True
            if row_touched:
                total_rows += 1

    # 拼接输出路径：原文件夹下的 output 子文件夹，文件名加“_已换算”
    base = os.path.basename(path)           # 取文件名（含扩展名）
    name, ext = os.path.splitext(base)      # 拆成“名字”和“.xlsx”
    out_dir = os.path.join(os.path.dirname(path), 'output')  # 原目录/output
    os.makedirs(out_dir, exist_ok=True)     # 不存在就建，存在也不报错
    out_path = os.path.join(out_dir, f'{name}_已换算{ext}')  # f-string 拼路径

    wb.save(out_path)                       # 另存为新文件，原文件不动
    return total_rows


def main():
    """程序入口：解析参数、遍历文件、逐个处理并打印结果"""
    parser = argparse.ArgumentParser(description='Excel 金额 元->万元 批量换算')
    parser.add_argument(
        '--folder',
        default=r'D:\wokbuddy2\AI自学\python练习\报表练习',  # 默认报表文件夹
        help='报表所在文件夹路径',
    )
    parser.add_argument(
        '-d', '--decimals',
        type=int, choices=[0, 1, 2], default=2,  # 只允许 0/1/2，默认 2
        help='保留小数位：0=整数, 1=1位, 2=2位（默认）',
    )
    args = parser.parse_args()              # 解析命令行参数

    if not os.path.isdir(args.folder):      # 文件夹不存在就报错退出
        print(f'文件夹不存在: {args.folder}')
        sys.exit(1)

    # 只处理 .xlsx 和 .xlsm（openpyxl 不支持老版 .xls）
    files = [f for f in os.listdir(args.folder)
             if f.lower().endswith(('.xlsx', '.xlsm'))]
    if not files:
        print('没找到可处理的 .xlsx / .xlsm 文件')
        return

    print(f'开始处理 | 小数位={args.decimals} | 共 {len(files)} 个文件')
    for f in files:
        path = os.path.join(args.folder, f)
        t0 = time.perf_counter()            # 记录开始时刻（高精度计时）
        try:
            rows = convert_file(path, args.decimals)
            dt = time.perf_counter() - t0   # 耗时 = 现在 - 开始
            # 成功：打印 文件名 + 处理行数 + 耗时
            print(f'[OK]   {f} | 处理行数={rows} | 耗时={dt:.2f}s')
        except Exception as e:              # 捕获任何异常，不中断整体
            dt = time.perf_counter() - t0
            # 失败：打印错误信息，继续下一个文件
            print(f'[FAIL] {f} | 错误={e} | 耗时={dt:.2f}s')


# 只有直接运行本文件时才执行 main()；被 import 时不执行
if __name__ == '__main__':
    main()
